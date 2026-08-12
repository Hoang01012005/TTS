"""
Script tạo slide PowerPoint báo cáo dự án VITS TTS Tiếng Việt.
Chạy: python create_slides.py
Output: output/VITS_TTS_BaoCao.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
import os, sys

# Fix encoding cho Windows console
if sys.platform == "win32":
    os.environ["PYTHONUTF8"] = "1"
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# ──────────────────────── Cấu hình ────────────────────────
OUTPUT_DIR = Path(__file__).resolve().parent / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_DIR / "VITS_TTS_BaoCao.pptx"

# ──────────────────────── Màu sắc ────────────────────────
BG_DARK = RGBColor(0x0F, 0x17, 0x2A)       # Nền tối xanh đen
BG_CARD = RGBColor(0x1A, 0x25, 0x3C)       # Nền card
ACCENT_BLUE = RGBColor(0x38, 0xBD, 0xF8)   # Xanh dương sáng
ACCENT_GREEN = RGBColor(0x4A, 0xDE, 0x80)  # Xanh lá
ACCENT_PURPLE = RGBColor(0xA7, 0x8B, 0xFA) # Tím
ACCENT_ORANGE = RGBColor(0xFB, 0x92, 0x3C) # Cam
ACCENT_RED = RGBColor(0xF8, 0x71, 0x71)    # Đỏ nhạt
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xB8, 0xC8)
MID_GRAY = RGBColor(0x64, 0x74, 0x8B)


# ──────────────────────── Hàm tiện ích ────────────────────
def set_slide_bg(slide, color):
    """Đặt màu nền cho slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape_bg(slide, left, top, width, height, color, alpha=None):
    """Thêm hình chữ nhật làm nền."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    # Rounded corners
    shape.adjustments[0] = 0.05
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 font_color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name="Segoe UI"):
    """Thêm text box vào slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_text(slide, left, top, width, height, items, font_size=16,
                    font_color=LIGHT_GRAY, bullet_color=ACCENT_BLUE):
    """Thêm danh sách bullet points."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"▸  {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.name = "Segoe UI"
        p.space_after = Pt(8)
    return txBox


def add_table(slide, left, top, width, height, rows, cols, data,
              header_color=ACCENT_BLUE):
    """Thêm bảng dữ liệu."""
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    # Style header
    for c in range(cols):
        cell = table.cell(0, c)
        cell.text = data[0][c]
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.name = "Segoe UI"
            p.alignment = PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    # Style data rows
    for r in range(1, rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.color.rgb = LIGHT_GRAY
                p.font.name = "Segoe UI"
                p.alignment = PP_ALIGN.CENTER
            if r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x15, 0x20, 0x35)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = BG_CARD

    return table_shape


def add_accent_line(slide, left, top, width, color=ACCENT_BLUE):
    """Thêm đường kẻ accent."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


# ──────────────────────── Tạo Slides ─────────────────────

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW = prs.slide_width
    SH = prs.slide_height

    # ═══════════════════ SLIDE 1: TRANG BÌA ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    set_slide_bg(slide, BG_DARK)

    # Gradient accent shape
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.08))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()

    # Title
    add_text_box(slide, Inches(1.5), Inches(1.8), Inches(10), Inches(1),
                 "🎤  VITS — Text-to-Speech", 44, ACCENT_BLUE, True)
    add_text_box(slide, Inches(1.5), Inches(2.8), Inches(10), Inches(0.8),
                 "Tổng hợp giọng nói tiếng Việt bằng mô hình VITS", 28, WHITE)

    add_accent_line(slide, Inches(1.5), Inches(3.8), Inches(2.5), ACCENT_BLUE)

    # Subtitle info
    add_text_box(slide, Inches(1.5), Inches(4.2), Inches(10), Inches(0.5),
                 "Variational Inference with adversarial learning for end-to-end Text-to-Speech",
                 18, LIGHT_GRAY)

    # Info boxes
    info_items = [
        ("📋  Framework", "Piper TTS + ONNX Runtime"),
        ("🗣️  Ngôn ngữ", "Tiếng Việt (vi)"),
        ("📊  Dataset", "VIVOS SPK01"),
        ("⚙️  Model", "VITS (epoch 4988)"),
    ]
    for i, (label, value) in enumerate(info_items):
        x = Inches(1.5 + i * 2.7)
        y = Inches(5.2)
        add_shape_bg(slide, x, y, Inches(2.4), Inches(1.2), BG_CARD)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.15), Inches(2), Inches(0.4),
                     label, 13, MID_GRAY)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.55), Inches(2), Inches(0.5),
                     value, 15, WHITE, True)

    # ═══════════════════ SLIDE 2: MỤC LỤC ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(1), Inches(0.5), Inches(5), Inches(0.8),
                 "📑  Nội dung báo cáo", 32, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(1), Inches(1.3), Inches(3))

    toc_items = [
        ("01", "Giới thiệu VITS", "Tổng quan về mô hình VITS và ứng dụng TTS"),
        ("02", "Kiến trúc mô hình", "Cấu trúc mạng, các thành phần chính"),
        ("03", "Dữ liệu huấn luyện", "Dataset VIVOS, tiền xử lý dữ liệu"),
        ("04", "Quá trình huấn luyện", "Cấu hình training, hyperparameters"),
        ("05", "Kết quả Inference", "Đánh giá chất lượng, tốc độ tổng hợp"),
        ("06", "Demo & Kết luận", "Demo thực tế và hướng cải thiện"),
    ]

    for i, (num, title, desc) in enumerate(toc_items):
        y = Inches(1.8 + i * 0.85)
        # Number circle
        add_shape_bg(slide, Inches(1), y, Inches(0.65), Inches(0.65), ACCENT_BLUE)
        add_text_box(slide, Inches(1), y + Inches(0.08), Inches(0.65), Inches(0.5),
                     num, 20, BG_DARK, True, PP_ALIGN.CENTER)
        # Title & desc
        add_text_box(slide, Inches(2), y + Inches(0.02), Inches(4), Inches(0.35),
                     title, 20, WHITE, True)
        add_text_box(slide, Inches(2), y + Inches(0.35), Inches(8), Inches(0.3),
                     desc, 14, MID_GRAY)

    # ═══════════════════ SLIDE 3: GIỚI THIỆU VITS ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.7),
                 "01 — Giới thiệu VITS", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # What is VITS
    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.2), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.7), Inches(5), Inches(0.5),
                 "VITS là gì?", 22, ACCENT_GREEN, True)
    add_bullet_text(slide, Inches(1.1), Inches(2.3), Inches(5), Inches(3.5), [
        "Variational Inference with adversarial learning\nfor end-to-end Text-to-Speech",
        "Kết hợp VAE + Normalizing Flow + GAN",
        "Tổng hợp giọng nói end-to-end (text → audio)",
        "Không cần vocoder riêng biệt (khác FastSpeech2)",
        "Chất lượng cao, tốc độ nhanh, tự nhiên",
        "Paper: Kim et al., ICML 2021",
    ], 15)

    # Why VITS
    add_shape_bg(slide, Inches(6.8), Inches(1.5), Inches(5.5), Inches(5.2), BG_CARD)
    add_text_box(slide, Inches(7.1), Inches(1.7), Inches(5), Inches(0.5),
                 "Tại sao chọn VITS?", 22, ACCENT_ORANGE, True)
    add_bullet_text(slide, Inches(7.1), Inches(2.3), Inches(5), Inches(3.5), [
        "End-to-end: Đơn giản hóa pipeline TTS",
        "Chất lượng giọng nói tự nhiên hơn Tacotron2",
        "Hỗ trợ multi-speaker (nhiều giọng nói)",
        "Tốc độ inference nhanh (real-time)",
        "Dễ export sang ONNX để deploy",
        "Cộng đồng lớn, nhiều công cụ hỗ trợ (Piper)",
    ], 15)

    # ═══════════════════ SLIDE 4: KIẾN TRÚC MÔ HÌNH ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.7),
                 "02 — Kiến trúc mô hình VITS", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # Architecture components
    components = [
        ("🔤 Text Encoder", ACCENT_BLUE,
         "Transformer Encoder\n• 6 layers, 2 heads\n• Hidden: 192 channels\n• Filter: 768 channels\n• Kernel size: 3\n• Dropout: 0.1"),
        ("🔗 Stochastic Duration\nPredictor", ACCENT_PURPLE,
         "Dự đoán thời lượng\n• Flow-based\n• Noise scale: 0.667\n• Noise W: 0.8\n• Length scale: 1.0"),
        ("🌊 Posterior Encoder\n+ Flow", ACCENT_GREEN,
         "VAE + Normalizing Flow\n• Inter channels: 192\n• Q Layers: 3\n• Variational inference\n• Latent representation"),
        ("🔊 HiFi-GAN Decoder", ACCENT_ORANGE,
         "Vocoder tích hợp\n• Upsample: [8,8,2,2]\n• ResBlock kernels: [3,7,11]\n• Dilation: [[1,3,5]×3]\n• Init channel: 512"),
    ]

    for i, (title, color, desc) in enumerate(components):
        x = Inches(0.6 + i * 3.1)
        y = Inches(1.5)
        add_shape_bg(slide, x, y, Inches(2.9), Inches(5.2), BG_CARD)
        # Color accent bar on top
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(2.9), Pt(4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        # Title
        add_text_box(slide, x + Inches(0.2), y + Inches(0.3), Inches(2.5), Inches(0.8),
                     title, 17, color, True)
        # Description
        add_text_box(slide, x + Inches(0.2), y + Inches(1.3), Inches(2.5), Inches(3.5),
                     desc, 13, LIGHT_GRAY)

    # Num symbols info
    add_text_box(slide, Inches(0.8), Inches(6.85), Inches(12), Inches(0.4),
                 "📊 Tổng số symbols: 256  |  Phoneme type: espeak  |  Sample rate: 22050 Hz  |  Single speaker",
                 14, MID_GRAY, alignment=PP_ALIGN.CENTER)

    # ═══════════════════ SLIDE 5: DỮ LIỆU HUẤN LUYỆN ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.7),
                 "03 — Dữ liệu huấn luyện", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # Dataset info card
    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.3), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.7), Inches(5), Inches(0.5),
                 "📂  Dataset: VIVOS SPK01", 22, ACCENT_GREEN, True)

    data_table = [
        ["Thông số", "Giá trị"],
        ["Nguồn dữ liệu", "VIVOS (AILab HCMUS)"],
        ["Speaker", "VIVOSSPK01 (single)"],
        ["Số lượng câu", "250 câu"],
        ["Tổng thời lượng", "~10 phút"],
        ["Sample rate", "22050 Hz"],
        ["Định dạng", "WAV 16-bit mono"],
        ["Text cleaners", "vietnamese_cleaners"],
        ["Add blank", "True"],
        ["Train / Val split", "240 / 10 câu"],
    ]
    add_table(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4),
              len(data_table), 2, data_table)

    # Processing pipeline card
    add_shape_bg(slide, Inches(7.0), Inches(1.5), Inches(5.5), Inches(5.3), BG_CARD)
    add_text_box(slide, Inches(7.3), Inches(1.7), Inches(5), Inches(0.5),
                 "⚙️  Pipeline xử lý dữ liệu", 22, ACCENT_PURPLE, True)

    pipeline_steps = [
        "1. Thu thập audio WAV từ VIVOS corpus",
        "2. Chuẩn hóa text (vietnamese_cleaners)",
        "3. Phonemize bằng espeak-ng (voice: vi)",
        "4. Tạo phoneme_id_map (156 ký hiệu IPA)",
        "5. Trích xuất mel-spectrogram:",
        "     • n_mel_channels = 80",
        "     • hop_length = 256",
        "     • win_length = 1024",
        "6. Chia train/val (metadata.csv)",
        "7. Export sang định dạng Piper",
    ]
    add_bullet_text(slide, Inches(7.3), Inches(2.4), Inches(5), Inches(4),
                    pipeline_steps, 14)

    # ═══════════════════ SLIDE 6: QUÁ TRÌNH HUẤN LUYỆN ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.7),
                 "04 — Quá trình huấn luyện", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # Training config card
    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.3), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.7), Inches(5), Inches(0.5),
                 "🏋️  Cấu hình Training", 22, ACCENT_GREEN, True)

    train_table = [
        ["Hyperparameter", "Giá trị"],
        ["Epochs", "10,000 (đã train: 4,988)"],
        ["Batch size", "16"],
        ["Learning rate", "2e-4"],
        ["LR decay", "0.99988"],
        ["Betas", "[0.8, 0.99]"],
        ["Epsilon", "1e-9"],
        ["Segment size", "8192"],
        ["FP16", "True"],
        ["c_mel / c_kl", "45 / 1.0"],
    ]
    add_table(slide, Inches(1.0), Inches(2.4), Inches(5.2), Inches(4),
              len(train_table), 2, train_table)

    # Training notes
    add_shape_bg(slide, Inches(7.0), Inches(1.5), Inches(5.5), Inches(5.3), BG_CARD)
    add_text_box(slide, Inches(7.3), Inches(1.7), Inches(5), Inches(0.5),
                 "📝  Ghi chú huấn luyện", 22, ACCENT_ORANGE, True)

    add_bullet_text(slide, Inches(7.3), Inches(2.4), Inches(5), Inches(4), [
        "Framework: Piper TTS (dựa trên PyTorch Lightning)",
        "Model đã train được 4,988 / 10,000 epochs",
        "Export ONNX: model_epoch_4988.onnx (~60MB)",
        "Optimizer: AdamW với LR scheduling",
        "Loss function:",
        "    • Mel-spectrogram loss (c_mel=45)",
        "    • KL divergence loss (c_kl=1.0)",
        "    • Adversarial loss (GAN)",
        "Mixed precision training (FP16) enabled",
        "Checkpoint: Generator (G) + Discriminator (D)",
    ], 14)

    # ═══════════════════ SLIDE 7: KẾT QUẢ INFERENCE ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
                 "05 — Kết quả Inference", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # Results table
    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(11.5), Inches(3.5), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.6), Inches(5), Inches(0.5),
                 "📊  Bảng kết quả tổng hợp giọng nói", 20, ACCENT_GREEN, True)

    results_table = [
        ["#", "Câu test", "Inference", "Audio", "RTF"],
        ["1", "Xin chào, tôi là trợ lý ảo tiếng Việt.", "0.14s", "2.47s", "0.056"],
        ["2", "Hôm nay thời tiết rất đẹp, chúng ta hãy đi dạo nhé.", "0.12s", "3.07s", "0.040"],
        ["3", "Trí tuệ nhân tạo đang thay đổi thế giới.", "0.10s", "2.44s", "0.041"],
        ["4", "Việt Nam là đất nước tươi đẹp với nhiều danh lam...", "0.14s", "3.18s", "0.043"],
        ["5", "Cảm ơn bạn đã sử dụng mô hình tổng hợp giọng nói.", "0.12s", "3.00s", "0.042"],
    ]
    add_table(slide, Inches(1.0), Inches(2.2), Inches(11), Inches(2.5),
              len(results_table), 5, results_table)

    # Stats cards
    stats = [
        ("⚡ Avg Inference", "0.12s", ACCENT_BLUE),
        ("🔊 Avg Audio", "2.83s", ACCENT_GREEN),
        ("📈 Avg RTF", "0.044", ACCENT_PURPLE),
        ("🖥️ Platform", "CPU (ONNX Runtime)", ACCENT_ORANGE),
    ]
    for i, (label, value, color) in enumerate(stats):
        x = Inches(0.8 + i * 3.0)
        y = Inches(5.3)
        add_shape_bg(slide, x, y, Inches(2.7), Inches(1.5), BG_CARD)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.15), Inches(2.3), Inches(0.4),
                     label, 14, MID_GRAY)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.6), Inches(2.3), Inches(0.6),
                     value, 28, color, True)

    # ═══════════════════ SLIDE 8: LONG TEXT TEST ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
                 "05b — Test đoạn văn bản dài", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(11.5), Inches(2.8), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.7), Inches(11), Inches(0.4),
                 "📝 Đoạn văn test (~100 từ, 4 câu):", 16, MID_GRAY)
    add_text_box(slide, Inches(1.1), Inches(2.2), Inches(11), Inches(1.8),
                 "\"Việt Nam là một quốc gia nằm ở phía đông bán đảo Đông Dương, thuộc khu vực "
                 "Đông Nam Á. Với đường bờ biển dài hơn ba nghìn hai trăm kilômét, Việt Nam sở hữu "
                 "nhiều bãi biển đẹp và vịnh nổi tiếng thế giới như vịnh Hạ Long, đã được UNESCO "
                 "công nhận là di sản thiên nhiên thế giới. Đất nước hình chữ S này có nền văn hóa "
                 "lâu đời với hơn bốn nghìn năm lịch sử...\"",
                 14, LIGHT_GRAY)

    # Long text results
    long_stats = [
        ("⏱️ Thời gian Inference", "1.24 giây", ACCENT_BLUE),
        ("🔊 Thời lượng Audio", "26.81 giây", ACCENT_GREEN),
        ("📈 Real-Time Factor", "0.046", ACCENT_PURPLE),
        ("🚀 Nhanh hơn real-time", "~22 lần", ACCENT_ORANGE),
    ]
    for i, (label, value, color) in enumerate(long_stats):
        x = Inches(0.8 + i * 3.0)
        y = Inches(4.7)
        add_shape_bg(slide, x, y, Inches(2.7), Inches(2.2), BG_CARD)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.2), Inches(2.3), Inches(0.5),
                     label, 14, MID_GRAY)
        add_text_box(slide, x + Inches(0.2), y + Inches(0.8), Inches(2.3), Inches(0.7),
                     value, 32, color, True)
        add_text_box(slide, x + Inches(0.2), y + Inches(1.6), Inches(2.3), Inches(0.4),
                     "✅ Real-time capable", 12, ACCENT_GREEN)

    # ═══════════════════ SLIDE 9: DEMO ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
                 "06 — Demo & Cách sử dụng", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # Pipeline
    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(11.5), Inches(1.5), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.6), Inches(5), Inches(0.4),
                 "🔄  Pipeline TTS", 18, ACCENT_GREEN, True)

    pipeline_boxes = [
        ("📝 Text Input", ACCENT_BLUE),
        ("🔤 Phonemize\n(espeak-ng)", ACCENT_PURPLE),
        ("🔢 Phoneme IDs", ACCENT_GREEN),
        ("🧠 VITS ONNX\nInference", ACCENT_ORANGE),
        ("🔊 Audio WAV", ACCENT_RED),
    ]
    for i, (label, color) in enumerate(pipeline_boxes):
        x = Inches(1.0 + i * 2.2)
        y = Inches(2.1)
        add_shape_bg(slide, x, y, Inches(1.8), Inches(0.7), color)
        add_text_box(slide, x + Inches(0.05), y + Inches(0.05), Inches(1.7), Inches(0.6),
                     label, 11, BG_DARK, True, PP_ALIGN.CENTER)
        if i < 4:
            add_text_box(slide, x + Inches(1.85), y + Inches(0.15), Inches(0.3), Inches(0.4),
                         "→", 20, MID_GRAY, True, PP_ALIGN.CENTER)

    # Usage
    add_shape_bg(slide, Inches(0.8), Inches(3.3), Inches(5.5), Inches(3.7), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(3.5), Inches(5), Inches(0.4),
                 "💻  Cách chạy", 18, ACCENT_BLUE, True)
    add_text_box(slide, Inches(1.1), Inches(4.0), Inches(5), Inches(2.8),
                 "# Cài đặt\n"
                 "pip install onnxruntime numpy\n"
                 "pip install piper-phonemize-fix\n\n"
                 "# Chạy 5 câu mẫu\n"
                 "python test_model.py\n\n"
                 "# Chạy câu tùy chọn\n"
                 "python test_model.py \"Xin chào\"\n\n"
                 "# Chế độ interactive\n"
                 "python test_model.py --interactive",
                 12, ACCENT_GREEN, font_name="Consolas")

    # Requirements
    add_shape_bg(slide, Inches(6.7), Inches(3.3), Inches(5.6), Inches(3.7), BG_CARD)
    add_text_box(slide, Inches(7.0), Inches(3.5), Inches(5), Inches(0.4),
                 "📋  Yêu cầu hệ thống", 18, ACCENT_ORANGE, True)
    add_bullet_text(slide, Inches(7.0), Inches(4.0), Inches(5), Inches(2.5), [
        "Python 3.11+ (khuyến nghị 3.14)",
        "ONNX Runtime 1.28+",
        "piper-phonemize-fix (Windows wheel)",
        "NumPy, SoundFile",
        "GPU: NVIDIA (optional, cần CUDA 13+)",
        "RAM: ~500MB cho model inference",
        "Disk: ~60MB (ONNX model file)",
    ], 14)

    # ═══════════════════ SLIDE 10: HƯỚNG CẢI THIỆN & KẾT LUẬN ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
                 "06b — Hướng cải thiện & Kết luận", 30, ACCENT_BLUE, True)
    add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2.5))

    # Improvements
    add_shape_bg(slide, Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.3), BG_CARD)
    add_text_box(slide, Inches(1.1), Inches(1.7), Inches(5), Inches(0.5),
                 "🚀  Hướng cải thiện", 22, ACCENT_ORANGE, True)

    improvements = [
        "⭐ Tăng dữ liệu: 10 phút → 3-5 giờ audio",
        "⭐ Train thêm: 4,988 → 15,000-20,000 epochs",
        "📊 Đánh giá MOS (Mean Opinion Score)",
        "🎯 Thử multi-speaker training",
        "⚡ Tối ưu model: quantization, pruning",
        "🌐 Tích hợp API server (FastAPI/Flask)",
        "📱 Deploy edge device (mobile, embedded)",
        "🔧 Fine-tune learning rate scheduling",
    ]
    add_bullet_text(slide, Inches(1.1), Inches(2.3), Inches(5.2), Inches(4),
                    improvements, 14)

    # Conclusion
    add_shape_bg(slide, Inches(7.0), Inches(1.5), Inches(5.5), Inches(5.3), BG_CARD)
    add_text_box(slide, Inches(7.3), Inches(1.7), Inches(5), Inches(0.5),
                 "✅  Kết luận", 22, ACCENT_GREEN, True)
    add_bullet_text(slide, Inches(7.3), Inches(2.3), Inches(5), Inches(4), [
        "Đã xây dựng thành công mô hình VITS TTS\ncho tiếng Việt",
        "Export ONNX cho phép deploy linh hoạt\ntrên nhiều nền tảng",
        "Inference nhanh: RTF ~0.04\n(nhanh hơn real-time 22 lần)",
        "Hỗ trợ phonemize tiếng Việt chính xác\nvới espeak-ng",
        "Cơ sở tốt để phát triển thêm:\nthêm data, train epochs, multi-speaker",
    ], 14)

    # ═══════════════════ SLIDE 11: CẢM ƠN ═══════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_DARK)

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.08))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()

    add_text_box(slide, Inches(1.5), Inches(2.2), Inches(10), Inches(1),
                 "🙏  Cảm ơn đã lắng nghe!", 44, ACCENT_BLUE, True, PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1.5), Inches(3.5), Inches(10), Inches(0.7),
                 "Dự án VITS Text-to-Speech Tiếng Việt", 24, WHITE, False, PP_ALIGN.CENTER)

    add_accent_line(slide, Inches(5.5), Inches(4.4), Inches(2.5), ACCENT_BLUE)

    add_text_box(slide, Inches(1.5), Inches(5.0), Inches(10), Inches(0.5),
                 "Q & A — Câu hỏi và thảo luận", 20, LIGHT_GRAY, False, PP_ALIGN.CENTER)

    # Footer
    add_text_box(slide, Inches(1.5), Inches(6.2), Inches(10), Inches(0.4),
                 "Framework: Piper TTS  |  Model: VITS  |  Runtime: ONNX  |  Language: Vietnamese",
                 14, MID_GRAY, False, PP_ALIGN.CENTER)

    # ──────────────────── Lưu file ────────────────────
    prs.save(str(OUTPUT_FILE))
    print(f"✅ Đã tạo slide: {OUTPUT_FILE}")
    return OUTPUT_FILE


if __name__ == "__main__":
    create_presentation()
